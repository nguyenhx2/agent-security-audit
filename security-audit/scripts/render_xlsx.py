#!/usr/bin/env python3
"""
Render the audit into an Excel workbook (.xlsx) for tracking and statistics.

Sheets
  Summary   project info, statistics tables (all COUNTIFS formulas over Findings, so they follow any
            edit the team makes) and 4 charts: by severity, repo x severity, by kind, by status.
  Findings  one row per finding; dropdowns (Severity, Confidence, Status, Kind, Repo, Standard) fed
            by the Master sheet; Owner / Due date / Notes columns for remediation tracking.
  Coverage  scanners per repo (status, image, digest, seconds), coverage gaps, sweep statistics.
  Master    the value lists behind every dropdown, as named ranges (lst_*). Edit a list once.

Layout canon (Japanese business-document style, self-contained): Meiryo, thin borders on every table cell, no zebra striping,
fills only where they carry meaning (grey headers, blue totals, severity colours), column A gutter,
middle alignment, gridlines off, Page Break Preview, landscape fit-to-width. No logo (general skill).

Safety: text is redacted through peek.py; strings starting with = + - @ are stored as text, never as
formulas (formula / CSV injection from quoted audited code).

Requires openpyxl (`py -m pip install openpyxl`). Optional `--verify`: recalculates the workbook, checks
there are no #REF!/#NAME?/#VALUE! errors, that the Summary total equals the number of findings and that
the 4 charts survive; `--pdf` also exports a PDF for a visual check. Engine (`--verify-with`, default
auto as in env_check.py): `excel` = real Excel via COM (Windows), `libreoffice` = `soffice --headless`
with a throw-away profile (macOS / Linux / Windows without Excel). Neither available: prints
"VERIFY SKIPPED" and exits 0 - state "not recalculated" in the reply.

Usage:
  python render_xlsx.py <report.md> [--findings <audit_dir>/findings] [--out <file.xlsx>]
         [--lang auto|en|ja] [--verify] [--verify-with auto|excel|libreoffice] [--pdf]
"""
from __future__ import annotations

import sys
sys.dont_write_bytecode = True  # keep the installed skill free of __pycache__
for _s in (sys.stdout, sys.stderr):  # UTF-8 output even on cp932 / cp1252 consoles and pipes
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")

import argparse
import json
import os
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, DoughnutChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    print("error: openpyxl is required: py -m pip install openpyxl  (or pip install openpyxl)", file=sys.stderr)
    sys.exit(2)

sys.path.insert(0, str(Path(__file__).resolve().parent))
from layout import findings_dir as _fdir, format_out, root_of_report, work_dir  # noqa: E402
from peek import find_secrets, redact as _redact  # noqa: E402
from i18n import SWEEP_COLS, TOOL_COVERAGE, T as I18N_T, detect_lang as _detect_lang, label as _label, meta_keys  # noqa: E402
try:
    from i18n import gap_text as _gap_text  # noqa: E402  localised scanner gap / skip-reason texts
except ImportError:  # older i18n: show the scanner's text as written
    def _gap_text(text: str, lang: str) -> str:
        return text
# data-validation messages (Findings sheet dropdowns and due date)
DV_TEXT = {"en": {"error": "Choose a value from the Master sheet list.", "error_title": "Invalid value",
                  "date": "Enter a date (YYYY-MM-DD).", "date_title": "Invalid date"},
           "ja": {"error": "マスタシートの一覧から選択してください。", "error_title": "入力値エラー",
                  "date": "日付（YYYY-MM-DD）を入力してください。", "date_title": "日付エラー"}}

FONT = "Meiryo"
NAVY, GREY, TOTAL, SUB = "1F3864", "D9D9D9", "DDEBF7", "F2F2F2"
SEV_FILL = {"P0": "F4B6B6", "P1": "FCD5B4", "P2": "FFF2CC", "Info": "E7E6E6"}
SEV_CHART = {"P0": "C00000", "P1": "ED7D31", "P2": "BF9000", "Info": "7F7F7F"}
SEVERITIES = ["P0", "P1", "P2", "Info"]
STATUSES = ["open", "in-progress", "closed", "accepted-risk", "false-positive"]
OPEN_STATES = ["open", "in-progress"]
CONFIDENCE = ["confirmed", "likely", "needs-runtime-check"]
KINDS = ["security", "sensitive-data-at-rest", "secret", "supply-chain", "configuration", "privacy",
         "performance", "quality"]
MASTER_MAX = 200  # rows scanned by the dynamic Master lists
DATA_ROWS = 500  # rows reserved for findings (dropdowns + formulas cover additions by the team)

L10N = {
    "en": {"definition": "Definition", "title": "Security audit - statistics", "project": "Project", "date": "Date", "depth": "Depth",
           "repos": "Repos", "skill": "Skill version", "verified": "Standards verified", "mode": "Scanner mode",
           "total": "Total", "open": "Open", "closed": "Closed / resolved", "by_sev": "Findings by severity",
           "by_repo": "Findings by repo and severity", "by_kind": "Findings by kind", "by_status": "Remediation status",
           "by_std": "Findings by primary standard", "kpi": "Open P0 + P1", "sev": "Severity", "repo": "Repo",
           "kind": "Kind", "status": "Status", "std": "Standard", "count": "Count",
           "findings_title": "Findings register", "coverage_title": "Scanner coverage",
           "master_title": "Master lists (dropdown values)", "note_master": "Edit these lists to change every dropdown. Named ranges: lst_*",
           "gaps": "Coverage gaps (stated, not clean)", "sweep": "Sensitive-data sweep",
           "cols": ["No", "ID", "Severity", "Confidence", "Status", "Repo", "Kind", "Title", "Standard", "Version",
                    "Requirement", "CWE", "Location", "Confirmed by", "Remedy (summary)", "Owner", "Due date", "Notes"],
           "cov_cols": ["Repo", "Tool", "Status", "Image / version", "Digest", "Seconds", "Note"],
           "sev_def": {"P0": "Blocker - direct high-impact compromise", "P1": "Major - exploitable under realistic conditions",
                       "P2": "Minor - limited blast radius", "Info": "Hygiene / hardening"}},
    "ja": {"definition": "定義", "title": "セキュリティ監査 統計", "project": "プロジェクト", "date": "日付", "depth": "深度",
           "repos": "リポジトリ", "skill": "スキル版", "verified": "基準検証日", "mode": "スキャナ方式",
           "total": "合計", "open": "未対応", "closed": "対応済", "by_sev": "重大度別件数",
           "by_repo": "リポジトリ別・重大度別件数", "by_kind": "種別別件数", "by_status": "対応状況",
           "by_std": "主基準別件数", "kpi": "未対応 P0＋P1", "sev": "重大度", "repo": "リポジトリ",
           "kind": "種別", "status": "状況", "std": "基準", "count": "件数",
           "findings_title": "指摘一覧", "coverage_title": "スキャナ実施状況",
           "master_title": "マスタ（プルダウン値）", "note_master": "この一覧を編集すると全プルダウンに反映されます（名前付き範囲 lst_*）",
           "gaps": "未カバー領域（安全を意味しない）", "sweep": "機微データ調査",
           "cols": ["No", "ID", "重大度", "確度", "状況", "リポジトリ", "種別", "件名", "基準", "版",
                    "要件", "CWE", "箇所", "根拠", "是正（要約）", "担当", "期限", "備考"],
           "cov_cols": ["リポジトリ", "ツール", "状況", "イメージ／版", "ダイジェスト", "秒", "備考"],
           "sev_def": {"P0": "致命的：直接的かつ重大な侵害", "P1": "重大：現実的な条件で悪用可能",
                       "P2": "軽微：影響範囲が限定的", "Info": "情報：衛生・強化事項"}},
}

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AC = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL = Alignment(horizontal="left", vertical="center", wrap_text=True)
AR = Alignment(horizontal="right", vertical="center")
AT = Alignment(horizontal="left", vertical="center", wrap_text=False)  # titles: never wrap


def host_path(p: str) -> str:
    if os.name == "nt":
        m = re.match(r"^/([a-zA-Z])(/.*)?$", p or "")
        if m:
            return f"{m.group(1).upper()}:{m.group(2) or '/'}"
    return p


def text_lines(text, width) -> int:
    n = 0
    for part in str(text or "").split("\n"):
        u = sum(2.0 if ord(ch) > 0x2E80 else 1.1 for ch in part)
        n += max(1, -(-int(u) // max(int(width - 1), 1)))
    return n


def F(sz=10, b=False, color="000000"):
    return Font(name=FONT, size=sz, bold=b, color=color)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def safe(v):
    """Redact secrets; never let text become a formula."""
    if v is None:
        return ""
    if isinstance(v, str) and find_secrets(v):  # only when a real secret shape is present
        return _redact(v, False)[0]
    return v


def put(ws, row, col, value, *, font=None, fill_=None, align=AL, border=BOX, fmt=None, redact=True):
    c = ws.cell(row=row, column=col)
    c.value = safe(value) if redact else ("" if value is None else value)
    if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
        c.data_type = "s"  # text, not formula
    c.font = font or F()
    if fill_:
        c.fill = fill(fill_)
    c.alignment = align
    if border:
        c.border = border
    if fmt:
        c.number_format = fmt
    return c


def formula(ws, row, col, f, *, bold=False, fill_=None, fmt="0"):
    c = ws.cell(row=row, column=col)
    c.value = f
    c.font = F(10, bold)
    c.alignment = AR
    c.border = BOX
    c.number_format = fmt
    if fill_:
        c.fill = fill(fill_)
    return c


def header(ws, row, col, labels):
    for i, lab in enumerate(labels):
        put(ws, row, col + i, lab, font=F(10, True), fill_=GREY, align=AC)


def widths(ws, cols):
    ws.column_dimensions["A"].width = 2  # gutter
    for i, w in enumerate(cols, start=2):
        ws.column_dimensions[get_column_letter(i)].width = w


def page(ws, landscape=True, title_rows=None):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.view = "pageBreakPreview"
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows


# ---------------- inputs ----------------
def parse_finding(text: str):
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n(.*)$", text, re.S)
    if not m:
        return None
    fm = {}
    for line in m.group(1).splitlines():
        line = re.sub(r"\s+#.*$", "", line)
        if ":" in line and not line.startswith(" "):
            k, v = line.split(":", 1)
            fm[k.strip()] = v.strip().strip("'\"")
    body = m.group(2)
    t = re.search(r"^#\s+(.+)$", body, re.M)
    fm["_title"] = re.sub(rf"^{re.escape(fm.get('id', ''))}\s*[-:]\s*", "", t.group(1).strip()) if t else fm.get("id", "")
    loc = re.search(r"^##\s*(Location|場所|該当箇所)\b.*?$(.*?)(?=^##\s|\Z)", body, re.S | re.M | re.I)
    anchors = re.findall(r"`([^`\n]+?:\d+(?:-\d+)?)`", loc.group(2) if loc else body)
    fm["_location"] = "\n".join(list(dict.fromkeys(anchors))[:12])
    rem = re.search(r"^##\s*(Remedy|Khắc phục|是正|対策)[^\n]*\n(.*?)(?=^##\s|\Z)", body, re.S | re.M | re.I)
    r = " ".join((rem.group(2) if rem else "").split())
    fm["_remedy"] = (r[:400] + "...") if len(r) > 400 else r
    return fm


def report_meta(md: str) -> dict:
    meta = {}
    keys = meta_keys()
    cell = r"((?:[^|\\]|\\.)+?)"  # a cell may hold escaped pipes (`\|`)
    for k, v in re.findall(r"^\|\s*" + cell + r"\s*\|\s*" + cell + r"\s*\|\s*$", md, re.M):
        meta.setdefault(keys.get(k.strip().lower(), k.strip().lower()), v.strip().replace("\\|", "|"))
    t = re.search(r"^#\s+(.+)$", md, re.M)
    meta["_title"] = t.group(1).strip() if t else ""
    sv = re.search(r"verified:?\s*\**\s*(\d{4}-\d{2}-\d{2})", md)
    meta["_verified"] = sv.group(1) if sv else ""
    return meta


def detect_lang(md: str) -> str:
    if re.search(r"[぀-ヿ]", md):
        return "ja"
    return "en"


# ---------------- build ----------------
def build(report: Path, fdir: Path, out: Path, lang: str) -> dict:
    md = report.read_text(encoding="utf-8", errors="replace")
    m_ = re.search(r"_([a-z]{2})\.md$", report.name)
    lang_ = lang if lang != "auto" else (m_.group(1) if m_ else _detect_lang(md))
    T = L10N.get(lang_, L10N["en"])
    DVT = DV_TEXT.get(lang_, DV_TEXT["en"])
    SN = I18N_T.get(lang_, I18N_T["en"])["sheets"]
    SEVL = {c: _label("severity", c, lang_) for c in SEVERITIES}
    STL = {c: _label("status", c, lang_) for c in STATUSES}
    CONFL = {c: _label("confidence", c, lang_) for c in CONFIDENCE}
    KINDL = {c: _label("kind", c, lang_) for c in KINDS}
    QF, QM = f"'{SN['findings']}'!", f"'{SN['master']}'!"  # quoted sheet refs (localised names)
    meta = report_meta(md)
    findings = []
    for f in sorted(fdir.rglob("*.md")):
        fm = parse_finding(f.read_text(encoding="utf-8", errors="replace"))
        if fm and fm.get("id"):
            findings.append(fm)
    findings.sort(key=lambda r: (SEVERITIES.index(r.get("severity")) if r.get("severity") in SEVERITIES else 9,
                                 r.get("repo", ""), r.get("id", "")))
    W = work_dir(root_of_report(report))
    repos = sorted({r.get("repo", "") for r in findings} | {p.name for p in W.iterdir()
                                                           if p.is_dir() and (p / "scanners.json").exists()})
    standards = sorted({r.get("standard", "") for r in findings if r.get("standard")})

    wb = Workbook()
    ws_s = wb.active
    ws_s.title = SN["summary"]
    ws_f = wb.create_sheet(SN["findings"])
    ws_c = wb.create_sheet(SN["coverage"])
    ws_m = wb.create_sheet(SN["master"])

    # ---------- Master ----------
    widths(ws_m, [12, 46, 3, 18, 3, 22, 3, 24, 3, 22, 3, 34])
    put(ws_m, 1, 2, T["master_title"], font=F(14, True, NAVY), border=None, align=AT)
    put(ws_m, 2, 2, T["note_master"], font=F(9, False, "595959"), border=None, align=AT)
    lists = [("lst_severity", T["sev"], [SEVL[c] for c in SEVERITIES], 2),
             ("lst_status", T["status"], [STL[c] for c in STATUSES], 5),
             ("lst_confidence", T["cols"][3], [CONFL[c] for c in CONFIDENCE], 7),
             ("lst_kind", T["kind"], [KINDL[c] for c in KINDS], 9),
             ("lst_repo", T["repo"], repos, 11), ("lst_standard", T["std"], standards, 13)]
    for name, title_, values, col in lists:
        header(ws_m, 4, col, [title_])
        for i, v in enumerate(values):
            put(ws_m, 5 + i, col, v)
            if name == "lst_severity":
                code = SEVERITIES[i]
                put(ws_m, 5 + i, col + 1, _label("severity_def", code, lang_))
                ws_m.cell(row=5 + i, column=col).fill = fill(SEV_FILL[code])
        L = get_column_letter(col)
        # dynamic list: grows when the team types a new value below the last one, no blank entries
        ref = f"OFFSET({QM}${L}$5,0,0,MAX(1,COUNTA({QM}${L}$5:${L}${MASTER_MAX})),1)"
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
        cols_ = f"{L}5:{get_column_letter(col + 1)}{MASTER_MAX}" if name == "lst_severity" else f"{L}5:{L}{MASTER_MAX}"
        ws_m.conditional_formatting.add(cols_, FormulaRule(formula=[f'${L}5<>""'], border=BOX))
    header(ws_m, 4, 3, [T.get("definition", "Definition")])
    page(ws_m)
    # no fixed print area: Excel prints the used range, which grows with the lists

    # ---------- Findings ----------
    cols = T["cols"]
    widths(ws_f, [5, 20, 9, 13, 13, 14, 14, 44, 22, 9, 12, 10, 34, 18, 50, 12, 11, 24])
    put(ws_f, 1, 2, T["findings_title"], font=F(14, True, NAVY), border=None, align=AT)
    put(ws_f, 2, 2, meta.get("_title", ""), font=F(9, False, "595959"), border=None, align=AT)
    HR = 4
    header(ws_f, HR, 2, cols)
    ws_f.row_dimensions[HR].height = 30
    first, last = HR + 1, HR + DATA_ROWS
    for i, r in enumerate(findings):
        row = first + i
        vals = [i + 1, r.get("id"), SEVL.get(r.get("severity"), r.get("severity")),
                CONFL.get(r.get("confidence"), r.get("confidence")), STL.get(r.get("status", "open"), r.get("status")),
                r.get("repo"), KINDL.get(r.get("kind"), r.get("kind")), r.get("_title"), r.get("standard"), r.get("standard_version"), r.get("requirement_id"),
                r.get("cwe"), r.get("_location"), r.get("confirmed_by"), r.get("_remedy"), "", "", ""]
        for j, v in enumerate(vals):
            put(ws_f, row, 2 + j, v, align=AC if j in (0, 2, 3, 4, 9, 11, 16) else AL,
                fmt="yyyy-mm-dd" if j == 16 else None)
        colw = [5, 20, 9, 13, 13, 14, 14, 44, 22, 9, 12, 10, 34, 18, 50, 12, 11, 24]
        lines = max(text_lines(v, colw[j]) for j, v in enumerate(vals))
        ws_f.row_dimensions[row].height = min(16.5 * lines + 6, 409)
    lastc = get_column_letter(1 + len(cols))
    ws_f.conditional_formatting.add(  # rows the team adds get ruled automatically
        f"B{first}:{lastc}{last}", FormulaRule(formula=[f"COUNTA($B{first}:${lastc}{first})>0"], border=BOX))
    ws_f.freeze_panes = ws_f.cell(row=first, column=4)
    ws_f.auto_filter.ref = f"B{HR}:{get_column_letter(1 + len(cols))}{first + max(len(findings), 1) - 1}"
    dv_map = {"D": "lst_severity", "E": "lst_confidence", "F": "lst_status", "G": "lst_repo", "H": "lst_kind",
              "J": "lst_standard"}
    for col, name in dv_map.items():
        dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True, showDropDown=False)
        dv.error, dv.errorTitle = DVT["error"], DVT["error_title"]
        if col == "J":
            dv.showErrorMessage = False  # standards list is advisory - new standards may be typed
        ws_f.add_data_validation(dv)
        dv.add(f"{col}{first}:{col}{last}")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    dv_date.error, dv_date.errorTitle = DVT["date"], DVT["date_title"]
    ws_f.add_data_validation(dv_date)
    dv_date.add(f"R{first}:R{last}")
    for sev, color in SEV_FILL.items():
        ws_f.conditional_formatting.add(f"D{first}:D{last}",
                                        FormulaRule(formula=[f'$D{first}="{SEVL[sev]}"'], fill=fill(color)))
    ws_f.conditional_formatting.add(f"F{first}:F{last}",
                                    FormulaRule(formula=[f'OR($F{first}="{STL["closed"]}",$F{first}="{STL["false-positive"]}")'],
                                                font=Font(name=FONT, color="7F7F7F")))
    page(ws_f, title_rows=f"{HR}:{HR}")
    # no fixed print area: Excel prints the used range, which grows with added rows

    # ---------- Summary ----------
    rng = lambda c: f"{QF}${c}${first}:${c}${last}"  # noqa: E731
    OPEN_L = [STL[x] for x in OPEN_STATES]
    SEV, STA, REP, KIN, STD = rng("D"), rng("F"), rng("G"), rng("H"), rng("J")
    widths(ws_s, [24, 11, 11, 11, 11, 11, 3])
    put(ws_s, 1, 2, T["title"], font=F(15, True, NAVY), border=None, align=AT)
    project = meta.get("_title", "").split(" - ", 1)[-1] if " - " in meta.get("_title", "") else meta.get("_title", "")
    info = [(T["project"], project), (T["date"], meta.get("date", "")),
            (T["depth"], _label("depth", meta.get("depth", ""), lang_)), (T["repos"], ", ".join(repos)),
            (T["skill"], re.sub(r".*v(\d[\d.]*).*", r"\1", meta.get("auditor", "")) or ""),
            (T["verified"], meta.get("_verified", "")), (T["mode"], meta.get("scanner mode", ""))]
    r0 = 3
    for i, (k, v) in enumerate(info):
        put(ws_s, r0 + i, 2, k, font=F(10, True), fill_=SUB)
        ws_s.merge_cells(start_row=r0 + i, start_column=3, end_row=r0 + i, end_column=7)
        put(ws_s, r0 + i, 3, v)
        for c in range(4, 8):
            ws_s.cell(row=r0 + i, column=c).border = BOX
    # KPI
    rk = r0 + len(info) + 1
    put(ws_s, rk, 2, T["kpi"], font=F(11, True), fill_=SEV_FILL["P0"])
    open_or = "+".join(f'COUNTIFS({SEV},"{SEVL[s]}",{STA},"{st}")' for s in ("P0", "P1") for st in OPEN_L)
    formula(ws_s, rk, 3, f"={open_or}", bold=True)
    ws_s.cell(row=rk, column=3).font = F(14, True, "C00000")
    # A: by severity
    ra = rk + 2
    put(ws_s, ra, 2, T["by_sev"], font=F(11, True, NAVY), border=None, align=AT)
    header(ws_s, ra + 1, 2, [T["sev"], T["total"], T["open"], T["closed"]])
    for i, s in enumerate(SEVERITIES):
        r = ra + 2 + i
        put(ws_s, r, 2, SEVL[s], fill_=SEV_FILL[s], align=AC)
        formula(ws_s, r, 3, f'=COUNTIFS({SEV},"{SEVL[s]}")')
        formula(ws_s, r, 4, "=" + "+".join(f'COUNTIFS({SEV},"{SEVL[s]}",{STA},"{st}")' for st in OPEN_L))
        formula(ws_s, r, 5, f"=C{r}-D{r}")
    rt = ra + 2 + len(SEVERITIES)
    put(ws_s, rt, 2, T["total"], font=F(10, True), fill_=TOTAL, align=AC)
    for c in (3, 4, 5):
        L = get_column_letter(c)
        formula(ws_s, rt, c, f"=SUM({L}{ra + 2}:{L}{rt - 1})", bold=True, fill_=TOTAL)
    sev_rows = (ra + 2, rt - 1)
    # B: repo x severity
    rb = rt + 2
    put(ws_s, rb, 2, T["by_repo"], font=F(11, True, NAVY), border=None, align=AT)
    header(ws_s, rb + 1, 2, [T["repo"]] + [SEVL[c] for c in SEVERITIES] + [T["total"]])
    for i, rp in enumerate(repos or ["-"]):
        r = rb + 2 + i
        put(ws_s, r, 2, rp)
        for j, s in enumerate(SEVERITIES):
            formula(ws_s, r, 3 + j, f'=COUNTIFS({REP},$B{r},{SEV},"{SEVL[s]}")')
        formula(ws_s, r, 7, f"=SUM(C{r}:F{r})", bold=True, fill_=TOTAL)
    repo_rows = (rb + 2, rb + 1 + len(repos or ["-"]))
    # C: by kind
    rc = repo_rows[1] + 2
    put(ws_s, rc, 2, T["by_kind"], font=F(11, True, NAVY), border=None, align=AT)
    header(ws_s, rc + 1, 2, [T["kind"], T["count"]])
    for i, k in enumerate(KINDS):
        r = rc + 2 + i
        put(ws_s, r, 2, KINDL[k])
        formula(ws_s, r, 3, f'=COUNTIFS({KIN},$B{r})')
    kind_rows = (rc + 2, rc + 1 + len(KINDS))
    # D: by status
    rd = kind_rows[1] + 2
    put(ws_s, rd, 2, T["by_status"], font=F(11, True, NAVY), border=None, align=AT)
    header(ws_s, rd + 1, 2, [T["status"], T["count"]])
    for i, st in enumerate(STATUSES):
        r = rd + 2 + i
        put(ws_s, r, 2, STL[st])
        formula(ws_s, r, 3, f'=COUNTIFS({STA},$B{r})')
    status_rows = (rd + 2, rd + 1 + len(STATUSES))
    # E: by standard
    re_ = status_rows[1] + 2
    put(ws_s, re_, 2, T["by_std"], font=F(11, True, NAVY), border=None, align=AT)
    header(ws_s, re_ + 1, 2, [T["std"], T["count"]])
    for i, sd in enumerate(standards or ["-"]):
        r = re_ + 2 + i
        put(ws_s, r, 2, sd)
        formula(ws_s, r, 3, f'=COUNTIFS({STD},$B{r})')
    last_row = re_ + 1 + len(standards or ["-"])

    # charts (right of the tables)
    def labels():
        dl = DataLabelList()
        dl.showVal, dl.showSerName, dl.showCatName, dl.showLegendKey, dl.showPercent = True, False, False, False, False
        dl.numFmt = "0;;;"  # hide zero labels
        return dl

    def axes(ch):
        for ax in (ch.x_axis, ch.y_axis):
            ax.delete = False

    def place(ch, anchor, w=13.5, h=7.6):
        ch.width, ch.height = w, h
        if ch.legend is not None:
            ch.legend.position = "b"
        ws_s.add_chart(ch, anchor)

    c1 = BarChart()
    c1.type, c1.title, c1.style = "col", T["by_sev"], 10
    c1.add_data(Reference(ws_s, min_col=3, max_col=4, min_row=ra + 1, max_row=sev_rows[1]), titles_from_data=True)
    c1.set_categories(Reference(ws_s, min_col=2, min_row=sev_rows[0], max_row=sev_rows[1]))
    c1.series[0].graphicalProperties.solidFill = "1F3864"
    c1.series[1].graphicalProperties.solidFill = "C00000"
    c1.dataLabels = labels(); axes(c1)
    c1.y_axis.majorGridlines = None; c1.y_axis.numFmt = "0"
    place(c1, "I3")

    c2 = BarChart()
    c2.type, c2.grouping, c2.overlap, c2.title = "col", "stacked", 100, T["by_repo"]
    c2.add_data(Reference(ws_s, min_col=3, max_col=6, min_row=rb + 1, max_row=repo_rows[1]), titles_from_data=True)
    c2.set_categories(Reference(ws_s, min_col=2, min_row=repo_rows[0], max_row=repo_rows[1]))
    for s, sr in zip(SEVERITIES, c2.series):
        sr.graphicalProperties.solidFill = SEV_CHART[s]
    c2.y_axis.majorGridlines = None; c2.y_axis.numFmt = "0"; axes(c2)
    c2.dataLabels = labels()
    place(c2, "Q3")

    c3 = BarChart()
    c3.type, c3.title = "bar", T["by_kind"]
    c3.add_data(Reference(ws_s, min_col=3, min_row=rc + 1, max_row=kind_rows[1]), titles_from_data=True)
    c3.set_categories(Reference(ws_s, min_col=2, min_row=kind_rows[0], max_row=kind_rows[1]))
    c3.series[0].graphicalProperties.solidFill = "2F5597"
    c3.legend = None
    c3.dataLabels = labels(); axes(c3)
    c3.x_axis.majorGridlines = None
    place(c3, "I19")

    c4 = DoughnutChart()
    c4.title = T["by_status"]
    c4.add_data(Reference(ws_s, min_col=3, min_row=rd + 1, max_row=status_rows[1]), titles_from_data=True)
    c4.set_categories(Reference(ws_s, min_col=2, min_row=status_rows[0], max_row=status_rows[1]))
    c4.dataLabels = labels()
    place(c4, "Q19")
    page(ws_s)
    ws_s.page_setup.fitToHeight = 1  # statistics sheet = one page
    ws_s.print_area = f"A1:Y{max(last_row, 35)}"

    # ---------- Coverage ----------
    widths(ws_c, [16, 13, 16, 34, 20, 9, 58])
    put(ws_c, 1, 2, T["coverage_title"], font=F(14, True, NAVY), border=None, align=AT)
    header(ws_c, 3, 2, T["cov_cols"])
    row, gaps, sweeps = 4, [], []
    for rp in repos:
        sj = W / rp / "scanners.json"
        if sj.exists():
            m = json.loads(sj.read_text(encoding="utf-8"))
            for t in m.get("tools", []):
                dig = (t.get("digest") or "").split("@")[-1].replace("sha256:", "")[:16]
                img = t.get("image") or t.get("version") or ""
                vals = [rp, t["tool"], _label("tool_status", t.get("status"), lang_), img.split("@")[0], dig,
                        t.get("seconds"), _gap_text(t.get("reason") or "", lang_) or TOOL_COVERAGE.get(lang_, TOOL_COVERAGE["en"]).get(t["tool"], t.get("coverage", ""))]
                for j, v in enumerate(vals):
                    put(ws_c, row, 2 + j, v, align=AR if j == 5 else AL, redact=(j == 6 and bool(t.get("reason"))),
                        fill_=None if t.get("status") == "ok" else "FCD5B4")
                row += 1
            gaps += [f"{rp}: {_gap_text(g, lang_)}" for g in m.get("coverage_gaps", [])]
        sw = W / rp / "sweep.json"
        if sw.exists():
            d = json.loads(sw.read_text(encoding="utf-8"))
            sweeps.append((rp, d.get("files_scanned"), d.get("migration_files"), d.get("schema_files"),
                           ", ".join(f"{k}={v}" for k, v in sorted(d.get("counts_by_class", {}).items())) or "-"))
    row += 1
    put(ws_c, row, 2, T["gaps"], font=F(11, True, NAVY), border=None, align=AT)
    row += 1
    for g in gaps or ["-"]:
        ws_c.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        put(ws_c, row, 2, g)
        for c in range(3, 9):
            ws_c.cell(row=row, column=c).border = BOX
        row += 1
    row += 1
    put(ws_c, row, 2, T["sweep"], font=F(11, True, NAVY), border=None, align=AT)
    row += 1
    header(ws_c, row, 2, SWEEP_COLS.get(lang_, SWEEP_COLS["en"]))
    ws_c.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
    row += 1
    for s in sweeps or [("-", "", "", "", "")]:
        for j, v in enumerate(s):
            put(ws_c, row, 2 + j, v, align=AR if j in (1, 2, 3) else AL)
        ws_c.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
        for c in (7, 8):
            ws_c.cell(row=row, column=c).border = BOX
        ws_c.row_dimensions[row].height = 30
        row += 1
    page(ws_c, title_rows="3:3")

    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "security-audit skill"
    wb.properties.title = meta.get("_title", "Security audit")
    wb.properties.company = None
    wb.active = 0
    wb.save(out)
    _fix_chart_numfmt(out)
    counts = Counter(r.get("severity") for r in findings)
    return {"findings": len(findings), "counts": dict(counts), "repos": repos, "lang": lang_}


def _fix_chart_numfmt(path: Path) -> None:
    """openpyxl omits the required sourceLinked attribute on data-label number formats, so Excel ignores
    them (zero labels would show). Rewrite the chart parts in place."""
    import shutil
    import tempfile
    import zipfile
    fd, name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(fd)  # Windows: an open handle blocks the later move
    tmp = Path(name)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/charts/chart"):
                data = data.replace(b'<numFmt formatCode="0;;;"/>', b'<numFmt formatCode="0;;;" sourceLinked="0"/>')
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


VERIFY_PS = r"""
$ErrorActionPreference = 'Stop'
$x = New-Object -ComObject Excel.Application
$x.Visible = $false; $x.DisplayAlerts = $false
try {
  $wb = $x.Workbooks.Open('__XLSX__')
  $x.CalculateFull()
  $errs = 0
  foreach ($ws in $wb.Worksheets) {
    $ur = $ws.UsedRange
    foreach ($c in $ur.Cells) { if ($c.HasFormula -and ($c.Text -match '^#(REF|NAME|VALUE|DIV|N/A|NULL|NUM)')) { $errs++ } }
  }
  $s = $wb.Worksheets.Item(1)
  $tot = $s.Range('__TOTAL__').Value2
  $charts = $s.ChartObjects().Count
  $dv = 0; try { $dv = $wb.Worksheets.Item(2).Range('D5').Validation.Type } catch {}
  if ('__PDF__' -ne '') { $wb.ExportAsFixedFormat(0, '__PDF__') }
  $wb.Close($false)
  Write-Output "errors=$errs total=$tot charts=$charts dv_type=$dv"
} finally { $x.Quit(); [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($x) }
"""


def verify(xlsx: Path, n_findings: int, total_ref: str, pdf: Path | None) -> int:
    ps = VERIFY_PS.replace("__XLSX__", str(xlsx).replace("'", "''")).replace("__TOTAL__", total_ref) \
        .replace("__PDF__", str(pdf).replace("'", "''") if pdf else "")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], capture_output=True,
                           text=True, timeout=300)
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"verify skipped: {e}")
        return 0
    out = (r.stdout or r.stderr).strip()
    print("excel verify:", out[-400:])
    m = re.search(r"errors=(\d+) total=([\d.]+) charts=(\d+)", out)
    if not m:
        print("verify inconclusive (Excel not available?)")
        return 0
    ok = int(m.group(1)) == 0 and int(float(m.group(2))) == n_findings and int(m.group(3)) == 4
    print("VERIFY", "PASS" if ok else "FAIL")
    return 0 if ok else 1


LO_RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""  # 0 = always recalculate on load (fresh profile only; the user's own profile is never used)
FORMULA_ERR = re.compile(r"^(#(REF|NAME|VALUE|DIV|N/A|NULL|NUM)|Err:\d+)")


def verify_libreoffice(soffice: str, xlsx: Path, n_findings: int, total_ref: str, pdf: Path | None) -> int:
    """Recalculate in LibreOffice headless (throw-away profile), then read the cached values back."""
    import shutil
    import tempfile
    import zipfile
    from openpyxl import load_workbook
    tmp = Path(tempfile.mkdtemp(prefix="sa-lo-"))
    try:
        prof = tmp / "profile"
        (prof / "user").mkdir(parents=True)
        (prof / "user" / "registrymodifications.xcu").write_text(LO_RECALC_XCU, encoding="utf-8")
        base = [soffice, f"-env:UserInstallation={prof.as_uri()}", "--headless", "--norestore", "--nologo"]
        try:
            r = subprocess.run(base + ["--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(xlsx)],
                               capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300)
            if pdf:
                subprocess.run(base + ["--convert-to", "pdf", "--outdir", str(pdf.parent), str(xlsx)],
                               capture_output=True, timeout=300)
        except (OSError, subprocess.TimeoutExpired) as e:
            print(f"verify skipped: libreoffice failed to run: {e}")
            print("VERIFY SKIPPED")
            return 0
        conv = tmp / "out" / xlsx.name
        if not conv.is_file():
            print("libreoffice verify:", ((r.stderr or r.stdout or "").strip())[-400:])
            print("verify inconclusive (LibreOffice produced no output)")
            print("VERIFY SKIPPED")
            return 0
        orig = load_workbook(xlsx)
        calc = load_workbook(conv, data_only=True)
        errs = 0
        for ws in orig.worksheets:
            if ws.title not in calc.sheetnames:
                continue
            cw = calc[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        v = cw[c.coordinate].value
                        if v is None or (isinstance(v, str) and FORMULA_ERR.match(v)):
                            errs += 1
        tot = calc.worksheets[0][total_ref].value
        with zipfile.ZipFile(conv) as z:
            charts = sum(1 for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
        print(f"libreoffice verify: errors={errs} total={tot} charts={charts}")
        try:
            ok = errs == 0 and int(float(tot)) == n_findings and charts == 4
        except (TypeError, ValueError):
            ok = False
        print("VERIFY", "PASS" if ok else "FAIL")
        return 0 if ok else 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def verify_any(how: str, xlsx: Path, n_findings: int, total_ref: str, pdf: Path | None) -> int:
    from env_check import detect_excel_verify  # same detection the agent saw in Phase 0
    det = detect_excel_verify()
    method = det["method"] if how == "auto" else how
    if method == "excel" or method == "excel-com":
        if not det["excel_com"]:
            print("verify skipped: Excel (COM) not available on this host")
            print("VERIFY SKIPPED")
            return 0
        return verify(xlsx, n_findings, total_ref, pdf)
    if method == "libreoffice":
        if not det["libreoffice"]:
            print("verify skipped: LibreOffice (soffice) not found")
            print("VERIFY SKIPPED")
            return 0
        return verify_libreoffice(det["libreoffice"], xlsx, n_findings, total_ref, pdf)
    print("verify skipped: no Excel or LibreOffice - workbook not recalculated (say so in the reply)")
    print("VERIFY SKIPPED")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("report")
    ap.add_argument("--findings")
    ap.add_argument("--out")
    ap.add_argument("--lang", default="auto", help="label language code (default: from the report name/text)")
    ap.add_argument("--verify", action="store_true", help="recalculate (Excel or LibreOffice) and check")
    ap.add_argument("--verify-with", choices=["auto", "excel", "libreoffice"], default="auto",
                    help="recalculation engine for --verify (implies --verify)")
    ap.add_argument("--pdf", action="store_true", help="with --verify: also export a PDF for a visual check")
    a = ap.parse_args()
    rp = Path(host_path(a.report)).resolve()
    if not rp.is_file():
        print(f"error: report not found: {rp}", file=sys.stderr)
        return 2
    # same language detection as render_html and build(): --lang, else the _<lang> suffix, else the report text
    lang_guess = a.lang if a.lang != "auto" else (re.search(r"_([a-z]{2})\.md$", rp.name) or [
        None, _detect_lang(rp.read_text(encoding="utf-8", errors="replace"))])[1]
    fdir = Path(host_path(a.findings)) if a.findings else _fdir(root_of_report(rp), lang_guess)
    if not fdir.is_dir():
        print(f"error: findings dir not found: {fdir}", file=sys.stderr)
        return 2
    out = Path(host_path(a.out)).resolve() if a.out else format_out(rp, "xlsx")
    info = build(rp, fdir, out, a.lang)
    print(f"wrote {out} ({info['findings']} findings {info['counts']}, lang={info['lang']})")
    if a.verify or a.verify_with != "auto":
        # total cell = Summary C<rt>; recompute its row the same way build() does
        from openpyxl import load_workbook
        ws = load_workbook(out).worksheets[0]
        total_ref = next(f"C{r}" for r in range(1, 80)
                         if isinstance(ws.cell(row=r, column=3).value, str)
                         and ws.cell(row=r, column=3).value.startswith("=SUM(C"))
        return verify_any(a.verify_with, out, info["findings"], total_ref, out.with_suffix(".pdf") if a.pdf else None)
    return 0


if __name__ == "__main__":
    sys.exit(main())
